# -*- coding: utf-8 -*-
"""Inspect the workbooks produced by run-model-tests.js with an independent reader.

    python validate_export.py [OUT_DIR]

Checks the generated template-layout workbook and the filled Parameter_Template.xlsx
without using any code from the tool itself.
"""
import os
import sys
import zipfile
import tempfile

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

try:                       # the sheet names / section markers are non-ASCII
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(tempfile.gettempdir(), 'spec-param-tool')
fail = 0


def check(ok, label, detail=""):
    global fail
    print(("  PASS  " if ok else "  FAIL  ") + label + (("  -> " + str(detail)) if detail else ""))
    if not ok:
        fail += 1


def open_wb(name):
    path = os.path.join(OUT, name)
    if not os.path.exists(path):
        check(False, name + " exists", path)
        return None
    check(zipfile.ZipFile(path).testzip() is None, name + " is a valid zip")
    return openpyxl.load_workbook(path)


print("== generated template-layout workbook ==")
wb = open_wb("gen_param_sheet.xlsx")
if wb:
    check("Summary" in wb.sheetnames, "has Summary sheet", wb.sheetnames[:3])
    param_sheets = [n for n in wb.sheetnames if n != "Summary"]
    check(bool(param_sheets), "has parameter sheets", len(param_sheets))
    ws = wb[param_sheets[0]]
    check(ws.max_column >= 7, "7 template columns", ws.max_column)
    check(len(ws.merged_cells.ranges) > 5, "merged cells written", len(ws.merged_cells.ranges))
    row2 = [ws.cell(row=2, column=c).value for c in range(1, 8)]
    check(row2[:2] == ["채널", "RED"], "channel header row", row2)
    row3 = [ws.cell(row=3, column=c).value for c in range(1, 4)]
    check(row3[0] == "조명 축", "axis row present", row3)
    labels = [ws.cell(row=r, column=1).value for r in range(1, 40)]
    check("GV 밝기" in labels, "GV row present")
    check("영역" in labels, "area row present")
    check("파라미터" in labels, "parameter row present")
    # first numeric value cell of the first parameter row
    prow = next(r for r in range(1, 60) if ws.cell(row=r, column=1).value == "파라미터")
    vals = [ws.cell(row=prow, column=c).value for c in (3, 5, 7)]
    check(all(isinstance(v, (int, float)) for v in vals), "parameter values are numeric", vals)
    check(any(isinstance(ws.cell(row=r, column=3).value, (int, float))
              for r in range(4, 12)), "a GV value cell was written from user input")

print("== file 1: LightSpec + GV workbook ==")
wb_light = open_wb("gen_light.xlsx")
if wb_light:
    check("LightSpec" in wb_light.sheetnames and "LightSpec Grouped" in wb_light.sheetnames,
          "LightSpec listings present", wb_light.sheetnames)
    check(any(n not in ("LightSpec", "LightSpec Grouped") for n in wb_light.sheetnames),
          "parameter (GV) sheet present too", wb_light.sheetnames)
    wsl = wb_light["LightSpec"]
    check(wsl.max_row > 20, "LightSpec rows written", wsl.max_row)

print("== file 2: InspectionSpec workbook ==")
wb_insp = open_wb("gen_inspect.xlsx")
if wb_insp:
    check(wb_insp.sheetnames == ["InspectionSpec", "Comparison"],
          "exactly the InspectionSpec listing + comparison", wb_insp.sheetnames)
    check(wb_insp["InspectionSpec"].max_row > 20, "InspectionSpec rows written",
          wb_insp["InspectionSpec"].max_row)

print("== file 2 filtered to LIGHT2 (SR / NonMetal only) ==")
wb_l2 = open_wb("gen_inspect_light2.xlsx")
if wb_l2:
    check("InspectionSpec" in wb_l2.sheetnames, "light-filtered workbook opens", wb_l2.sheetnames)
    ws2 = wb_l2["InspectionSpec"]
    pns = {ws2.cell(row=r, column=2).value for r in range(1, ws2.max_row + 1)}
    check(any(str(v).startswith("▸") for v in pns if v), "keeps the section rows",
          [v for v in pns if v and str(v).startswith("▸")])
    check(all("NonMetal" in str(v) for v in pns if v and str(v).startswith("▸")),
          "only the NonMetal area is left", [v for v in pns if v and str(v).startswith("▸")])
    check(wb_insp is None or ws2.max_row < wb_insp["InspectionSpec"].max_row,
          "filtered sheet is shorter than the unfiltered one",
          "%s < %s" % (ws2.max_row, wb_insp["InspectionSpec"].max_row if wb_insp else "?"))

print("== filled real template ==")
wb2 = open_wb("filled_template.xlsx")
if wb2:
    check(wb2.sheetnames[:3] == ["조명지침", "영역 Convention", "DMG 조명 1번"], "original sheet order kept", wb2.sheetnames)
    ws = wb2["Top 조명 2번"] if "Top 조명 2번" in wb2.sheetnames else wb2[wb2.sheetnames[3]]
    nums = 0
    for row in ws.iter_rows(min_row=9, max_row=200, min_col=3, max_col=3):
        for c in row:
            if isinstance(c.value, (int, float)):
                nums += 1
    check(nums > 50, "value cells present in the template sheet", nums)
    ws_rules = wb2["조명지침"]
    check(isinstance(ws_rules["B2"].value, str) and len(ws_rules["B2"].value) > 5, "rules sheet untouched", ws_rules["B2"].value[:30])
    ws_tpl = wb2["Top 조명 3번"]
    check(any(isinstance(ws_tpl.cell(row=r, column=2).value, str) and "Defect" in str(ws_tpl.cell(row=r, column=2).value)
              for r in range(1, 120)), "parameter labels still in place")

print("== filled template + appended LightSpec listings ==")
wb3 = open_wb("filled_template_with_light.xlsx")
if wb3:
    check(wb3.sheetnames[:3] == ["조명지침", "영역 Convention", "DMG 조명 1번"],
          "template sheets still first", wb3.sheetnames[:4])
    check("LightSpec" in wb3.sheetnames and "LightSpec Grouped" in wb3.sheetnames,
          "LightSpec listings appended after the template", wb3.sheetnames[-2:])
    check(wb3["LightSpec"].max_row > 20, "appended LightSpec sheet has rows", wb3["LightSpec"].max_row)

print("\nVALIDATION " + ("FAILED" if fail else "OK"))
sys.exit(1 if fail else 0)
